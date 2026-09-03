from io import BytesIO

from docx import Document

from hireme_ai.schemas.candidate import CandidateProfile
from hireme_ai.schemas.resume import GeneratedResumeBullet


def render_docx(candidate: CandidateProfile, bullets: list[GeneratedResumeBullet]) -> bytes:
    doc = Document()
    doc.add_heading(candidate.name, 0)
    doc.add_heading("Verified Skill Extract", level=1)
    for bullet in bullets:
        doc.add_paragraph(bullet.text, style="List Bullet")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
